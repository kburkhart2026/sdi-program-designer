#!/usr/bin/env python3
"""
Merge the Tooling U course catalog into public/catalog.json.

    python3 scripts/merge_tooling_u_catalog.py ~/Downloads/ToolingU-SME-Catalog.xlsx

Tooling U publishes more courses than the SME certification mappings reference.
The mappings give us 283 credential-bearing lessons; the full catalog export has
639. This script adds the remainder under a synthetic credential, **Other**, so
instructional designers can pull non-credential content into a program without
it affecting any credential's coverage maths.

WHAT IT TOUCHES
  - Adds/replaces the `Other` credential.
  - Adds the unmapped courses to the deduplicated `classes` index with
    `creds: ["Other"]`.
  - Backfills `topic` and `description` onto lessons already in the file where
    the export knows them.
It never edits an SME credential's lesson list, and never changes an existing
lesson's `creds` or `url`.

IDEMPOTENT. The `Other` credential and every `Other`-only lesson are dropped and
rebuilt on each run, so re-running with a newer export removes withdrawn courses
instead of accumulating them. Run it, then eyeball `git diff public/catalog.json`.

`topic` is the export's FunctionalArea — the 11 headings the lesson rail groups
`Other` by. `dept` stays the real Tooling U Department so the program map's
"content mix by department" keeps working.

Requires: openpyxl.
"""

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  python3 -m pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
CATALOG = ROOT / "public" / "catalog.json"

OTHER_CODE = "Other"
OTHER_TITLE = "Tooling U catalog (no SME credential)"
SHEET = "Online Classes"


def load_export(xlsx: Path) -> dict[str, dict]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"expected a '{SHEET}' sheet; found {wb.sheetnames}")
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    required = {"ClassId", "ClassName", "Department", "FunctionalArea"}
    missing = required - set(header)
    if missing:
        sys.exit(f"export is missing column(s): {', '.join(sorted(missing))}")

    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        r = dict(zip(header, row))
        cid = str(r["ClassId"]).strip()
        if not cid:
            continue
        # First occurrence wins; the export has repeated a ClassId before.
        out.setdefault(cid, r)
    return out


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    xlsx = Path(sys.argv[1]).expanduser()
    if not xlsx.exists():
        sys.exit(f"no such file: {xlsx}")

    tu = load_export(xlsx)
    catalog = json.loads(CATALOG.read_text())

    # --- strip the previous run so this is a rebuild, not an append ----------
    catalog["credentials"] = [c for c in catalog["credentials"] if c["code"] != OTHER_CODE]
    catalog["classes"] = [
        c for c in catalog["classes"] if c.get("creds") != [OTHER_CODE]
    ]

    by_id = {c["id"]: c for c in catalog["classes"]}
    sme_ids = set(by_id)

    # --- backfill topic + description onto lessons we already carry ----------
    enriched = 0
    for cid, row in tu.items():
        existing = by_id.get(cid)
        if not existing:
            continue
        if row.get("FunctionalArea"):
            existing["topic"] = row["FunctionalArea"]
        if row.get("Description"):
            existing["description"] = str(row["Description"]).strip()
        enriched += 1

    # --- the unmapped remainder becomes `Other` ------------------------------
    unmapped = sorted(
        (cid for cid in tu if cid not in sme_ids),
        key=lambda i: (
            tu[i].get("FunctionalArea") or "",
            tu[i].get("Department") or "",
            tu[i].get("ClassName") or "",
        ),
    )

    other_classes = []
    for cid in unmapped:
        r = tu[cid]
        name = str(r["ClassName"]).strip()
        dept = str(r["Department"] or "Unassigned").strip()
        topic = str(r["FunctionalArea"] or "Unassigned").strip()
        desc = str(r["Description"]).strip() if r.get("Description") else ""

        # Per-credential list. `url` is "" because the export ships an empty
        # ClassLink column for every row; the deduplicated entry below records
        # that honestly as null, and the UI falls back to the Tooling U catalog.
        other_classes.append({"dept": dept, "name": name, "id": cid, "url": ""})

        entry = {
            "id": cid,
            "name": name,
            "dept": dept,
            "url": None,
            "creds": [OTHER_CODE],
            "topic": topic,
        }
        if desc:
            entry["description"] = desc
        catalog["classes"].append(entry)

    catalog["credentials"].append(
        {
            "code": OTHER_CODE,
            "title": OTHER_TITLE,
            "source": xlsx.name,
            "kind": "other",
            "classes": other_classes,
        }
    )

    catalog["classes"].sort(key=lambda c: c["id"])
    catalog["generated"] = catalog.get("generated", "")
    catalog["otherGenerated"] = __import__("datetime").date.today().isoformat()

    # Compact, matching the file's existing convention. It is generated data,
    # never hand-edited, and indentation would add ~75 KB for nobody's benefit.
    CATALOG.write_text(json.dumps(catalog, ensure_ascii=False, separators=(",", ":")))

    total = len(catalog["classes"])
    print(f"export            : {len(tu)} unique courses from {xlsx.name}")
    print(f"already mapped    : {enriched} (topic/description backfilled)")
    print(f"added as '{OTHER_CODE}'   : {len(other_classes)}")
    print(f"catalog total     : {total} lessons, {len(catalog['credentials'])} credentials")
    print(f"file size         : {CATALOG.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    main()
